"""Result exports: CSV, Excel and PDF."""
import csv
import io

from django.http import HttpResponse
from django.utils import timezone

from voting.views import sanitize_csv_value


def _rows(data):
    for position in data.get('positions', []):
        for candidate in position.get('candidates', []):
            yield [position.get('position'), position.get('ballot_type'), candidate.get('rank'), candidate.get('name'),
                   candidate.get('votes'), candidate.get('percentage'),
                   'yes' if str(candidate.get('candidate_id')) in map(str, position.get('winners', [])) else '']


HEADER = ['Position', 'Method', 'Rank', 'Candidate / option', 'Votes', 'Percent', 'Winner']


def _filename(event, ext):
    return f'results_{event.pk}_{timezone.now():%Y%m%d%H%M}.{ext}'


def results_csv(event, data):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{_filename(event, "csv")}"'
    writer = csv.writer(response)
    writer.writerow(HEADER)
    for row in _rows(data):
        writer.writerow([sanitize_csv_value(v) if isinstance(v, str) else v for v in row])
    turnout = data.get('turnout', {})
    writer.writerow([])
    writer.writerow(['Eligible voters', turnout.get('eligible_voters')])
    writer.writerow(['Votes cast', turnout.get('votes_cast')])
    writer.writerow(['Turnout %', turnout.get('turnout_percentage')])
    writer.writerow(['Invalid ballots', turnout.get('invalid_ballots')])
    return response


def results_xlsx(event, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Results'
    sheet.append([event.title])
    sheet['A1'].font = Font(bold=True, size=14)
    sheet.append([])
    sheet.append(HEADER)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    for row in _rows(data):
        sheet.append([sanitize_csv_value(v) if isinstance(v, str) else v for v in row])
    turnout = data.get('turnout', {})
    summary = workbook.create_sheet('Turnout')
    for label, key in (('Eligible voters', 'eligible_voters'), ('Votes cast', 'votes_cast'),
                       ('Turnout %', 'turnout_percentage'), ('Ballots counted', 'ballots_counted'),
                       ('Invalid ballots', 'invalid_ballots')):
        summary.append([label, turnout.get(key)])
    if data.get('constituencies'):
        breakdown = workbook.create_sheet('Constituencies')
        breakdown.append(['Constituency', 'Position', 'Candidate', 'Votes', 'Percent'])
        for entry in data['constituencies']:
            if entry.get('suppressed'):
                breakdown.append([entry['name'], '(suppressed: fewer ballots than the anonymity threshold)'])
                continue
            for position in entry.get('positions', []):
                for candidate in position.get('candidates', []):
                    breakdown.append([entry['name'], position.get('position'), candidate.get('name'),
                                      candidate.get('votes'), candidate.get('percentage')])
    for worksheet in workbook.worksheets:
        for column in worksheet.columns:
            width = max(len(str(c.value or '')) for c in column)
            worksheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{_filename(event, "xlsx")}"'
    return response


def results_pdf(event, data, certification=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, title=f'Results - {event.title}', leftMargin=1.8 * cm,
                                 rightMargin=1.8 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    wine = colors.HexColor('#800020')
    story = [Paragraph(f'<font color="#800020"><b>{_esc(event.title)}</b></font>', styles['Title']),
             Paragraph(f'Official results report &middot; generated {timezone.now():%d %b %Y %H:%M %Z}', styles['Normal']),
             Spacer(1, 10)]
    turnout = data.get('turnout', {})
    summary = [['Eligible voters', turnout.get('eligible_voters') if turnout.get('eligible_voters') is not None else '-'],
               ['Votes cast', turnout.get('votes_cast')],
               ['Turnout', f"{turnout.get('turnout_percentage')}%" if turnout.get('turnout_percentage') is not None else '-'],
               ['Ballots counted', turnout.get('ballots_counted')], ['Invalid ballots', turnout.get('invalid_ballots')]]
    table = Table(summary, colWidths=[6 * cm, 6 * cm])
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.25, colors.grey), ('BACKGROUND', (0, 0), (0, -1), colors.whitesmoke)]))
    story += [table, Spacer(1, 14)]
    for position in data.get('positions', []):
        story.append(Paragraph(f'<b>{_esc(position.get("position"))}</b> &middot; {_esc(position.get("method", ""))}', styles['Heading3']))
        rows = [['Rank', 'Candidate / option', 'Votes', '%', '']]
        winners = set(map(str, position.get('winners', [])))
        for candidate in position.get('candidates', []):
            rows.append([candidate.get('rank'), _esc(candidate.get('name')), candidate.get('votes'),
                         candidate.get('percentage'), 'WINNER' if str(candidate.get('candidate_id')) in winners else ''])
        table = Table(rows, colWidths=[1.5 * cm, 8 * cm, 2.5 * cm, 2 * cm, 2.5 * cm])
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), wine), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                                   ('GRID', (0, 0), (-1, -1), 0.25, colors.grey), ('FONTSIZE', (0, 0), (-1, -1), 9)]))
        story.append(table)
        extra = []
        if position.get('abstentions'):
            extra.append(f"Abstentions: {position['abstentions']}")
        if position.get('ties'):
            extra.append('TIE - requires resolution under the election rules')
        if position.get('passed') is not None:
            extra.append(f"Referendum {'PASSED' if position['passed'] else 'NOT PASSED'} (threshold {position.get('threshold')}%)")
        if extra:
            story.append(Paragraph(' &middot; '.join(extra), styles['Italic']))
        story.append(Spacer(1, 10))
    if certification is not None:
        story += [Spacer(1, 10), Paragraph('<b>Certification</b>', styles['Heading3']),
                  Paragraph(f'Certified {certification.certified_at:%d %b %Y %H:%M %Z} by '
                            f'{_esc(certification.certified_by.username if certification.certified_by else "-")}', styles['Normal']),
                  Paragraph(f'Result hash: <font face="Courier" size="8">{certification.payload["result_hash"]}</font>', styles['Normal']),
                  Paragraph(f'Bulletin board root: <font face="Courier" size="8">{certification.payload.get("bulletin_root", "")}</font>', styles['Normal']),
                  Paragraph(f'Ed25519 signature: <font face="Courier" size="7">{certification.signature}</font>', styles['Normal']),
                  Paragraph(f'Signing key id: {certification.key_id}', styles['Normal'])]
    document.build(story)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_filename(event, "pdf")}"'
    return response


def _esc(value):
    import html

    return html.escape(str(value if value is not None else ''), quote=False)


def export(event, data, fmt, certification=None):
    if fmt == 'csv':
        return results_csv(event, data)
    if fmt == 'xlsx':
        return results_xlsx(event, data)
    if fmt == 'pdf':
        return results_pdf(event, data, certification)
    raise ValueError('Unsupported format')
